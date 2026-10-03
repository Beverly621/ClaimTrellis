import hashlib
import io
import json
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from test_policy import choice, judgment
from test_revisions import RevisionProvider, reference
from test_typesafe_jev import response_body

from claim_trellis.api import create_app
from claim_trellis.config import Settings
from claim_trellis.deterministic import run_deterministic_checks
from claim_trellis.evidence_sets import evidence_text, make_evidence_set
from claim_trellis.ingestion import _clean_text, parse_document_bytes
from claim_trellis.models import AuditRequest, ProposalStatus, SourceAccessTier
from claim_trellis.policy import propose_disposition
from claim_trellis.provider import ProviderError
from claim_trellis.providers.typesafe_jev import TypeSafeJevProvider
from claim_trellis.retrieval import retrieve

REGRESSIONS = json.loads((Path(__file__).parents[1] / "benchmarks/p0_regression.json").read_text())[
    "cases"
]


@pytest.mark.parametrize("case", REGRESSIONS, ids=lambda case: case["id"])
def test_p0_engineering_regressions(case):
    checks = run_deterministic_checks(case["claim"], case["evidence"], case.get("quote"))
    assert any(item.blocking for item in checks.findings) == case["blocking"]
    if case["blocking"]:
        policy = propose_disposition(
            checks, judgment(), source_access_tier=SourceAccessTier.FULL_TEXT
        )
        assert policy.status != ProposalStatus.SUPPORTED
    assert checks.check_set_version == "deterministic-v2"


def assert_blocks(document):
    assert document.parser_version == "structured-document-v2"
    for block in document.blocks:
        assert document.text[block.char_start : block.char_end] == block.text
        assert hashlib.sha256(block.text.encode()).hexdigest() == block.sha256
    assert len({b.block_id for b in document.blocks}) == len(document.blocks)


def test_markdown_headings_and_txt_lines_preserve_text():
    raw = (
        "# Methods\r\n\r\n## Training\r\nIntervention used 20 mg.\n\n# Results\nOutcome decreased."
    )
    document = parse_document_bytes("source.md", raw.encode())
    assert document.text == _clean_text(raw)
    assert_blocks(document)
    training = next(b for b in document.blocks if "20 mg" in b.text)
    assert training.section == "Methods > Training"
    assert training.block_type == "paragraph"
    assert "heading:Methods > Training" in training.locator
    plain = parse_document_bytes("source.txt", b"First line.\nSecond line.\n\nFourth line.")
    assert_blocks(plain)
    assert plain.blocks[0].locator == "lines:1-2"
    assert plain.blocks[1].locator == "lines:4-4"


def test_docx_heading_paragraph_and_table_locators():
    doc = Document()
    doc.add_heading("Results", 1)
    doc.add_paragraph("The endpoint decreased.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "20 mg"
    data = io.BytesIO()
    doc.save(data)
    document = parse_document_bytes("source.docx", data.getvalue())
    assert document.text == "Results\n\nThe endpoint decreased.\n\n20 mg"
    assert_blocks(document)
    assert document.blocks[1].section == "Results"
    assert document.blocks[1].paragraph == 2
    assert document.blocks[-1].block_type == "table_row"


def test_pdf_uses_actual_page_numbers_without_inventing_sections():
    writer = PdfWriter()
    for text in ("Methods were documented.", "Outcome decreased."):
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    data = io.BytesIO()
    writer.write(data)
    document = parse_document_bytes("source.pdf", data.getvalue())
    assert_blocks(document)
    assert [b.page for b in document.blocks] == [1, 2]
    assert all(b.section is None for b in document.blocks)
    assert "page:2" in document.blocks[1].locator


def source_text():
    return "\n\n".join(
        f"Section {i}. The intervention decreased the endpoint in cohort {i}. "
        + "Unrelated methodological background. " * 22
        for i in range(12)
    )


def test_top_k_sets_are_grounded_bounded_and_deterministic():
    text = source_text()
    first = retrieve("The intervention decreased the endpoint in cohort 8.", text)
    assert len(first) == 10
    assert first == retrieve("The intervention decreased the endpoint in cohort 8.", text)
    for candidate in first:
        p = candidate.passage
        assert text[p.start_char : p.end_char] == p.text
        assert hashlib.sha256(p.text.encode()).hexdigest() == p.sha256
    subset = make_evidence_set([first[1].passage, first[0].passage])
    assert subset.sha256 == make_evidence_set(list(reversed(subset.passages))).sha256
    with pytest.raises(ValueError):
        make_evidence_set([first[0].passage] * 2)


def test_adjacent_limitations_are_preserved_in_context():
    text = "The intervention decreased the endpoint.\n\nOnly the secondary endpoint was measured."
    items = retrieve("The intervention decreased the endpoint.", text)
    assert "secondary endpoint" in items[0].passage.text
    assert len(items) == 1  # De-duplicated expanded source span.


def test_forged_document_blocks_are_rejected():
    document = parse_document_bytes("source.txt", b"The intervention decreased the endpoint.")
    forged = document.blocks[0].model_copy(update={"text": "Invented evidence."})
    with pytest.raises(ValueError, match="source-grounded"):
        AuditRequest(
            claim="The intervention decreased the endpoint.",
            source_text=document.text,
            source_blocks=[forged],
        )


def test_document_blocks_cannot_omit_non_whitespace_source():
    document = parse_document_bytes("source.txt", b"First paragraph.\n\nSecond paragraph.")
    with pytest.raises(ValueError, match="cover"):
        AuditRequest(
            claim="First paragraph.", source_text=document.text, source_blocks=document.blocks[:1]
        )


def test_source_hash_must_match_normalized_text():
    with pytest.raises(ValueError, match="content hash"):
        AuditRequest(
            claim="A claim.", source_text="Original source.", source={"content_sha256": "0" * 64}
        )


def test_numeric_matching_does_not_round_different_decimal_spellings():
    checks = run_deterministic_checks(
        "The value was 9007199254740993 mg.", "The value was 9007199254740992 mg."
    )
    assert checks.unmatched_claim_numbers


def test_full_text_metadata_does_not_prove_exhaustive_source_silence():
    policy = propose_disposition(
        run_deterministic_checks("The intervention helped.", "Unrelated selected passage."),
        judgment("not_addressed"),
        source_access_tier=SourceAccessTier.FULL_TEXT,
    )
    assert policy.status == ProposalStatus.REVIEW_REQUIRED
    assert "not exhaustive" in policy.reasons[0]


def test_ungrounded_format_locator_falls_back_without_omitting_text(monkeypatch):
    monkeypatch.setattr(
        "claim_trellis.ingestion._read_pdf",
        lambda data: (
            "First paragraph.\n\nSecond paragraph.",
            [],
            [("First paragraph.", {"page": 1})],
        ),
    )
    document = parse_document_bytes("source.pdf", b"test-only")
    assert_blocks(document)
    assert len(document.blocks) == 2
    assert all(block.page is None for block in document.blocks)
    assert any("exact text offsets" in warning for warning in document.warnings)


def test_block_manifest_does_not_persist_full_source_text(tmp_path):
    app = create_app(Settings(data_dir=tmp_path, jev_api_key=None))
    with TestClient(app) as client:
        source = source_text()
        audit = client.post(
            "/api/v1/audits",
            json={
                "claim": "The intervention decreased the endpoint.",
                "source_text": source,
                "use_judgment_provider": False,
            },
        ).json()
        assert audit["source_blocks"]
        assert all("text" not in block for block in audit["source_blocks"])
        events = client.get(f"/api/v1/audits/{audit['audit_id']}/events").json()
        parsed = next(event for event in events if event["event_type"] == "document.parsed")
        assert all("text" not in block for block in parsed["payload"]["blocks"])
        assert source not in json.dumps(audit)


@pytest.mark.parametrize(
    "field",
    [
        "outcome_alignment",
        "timeframe_alignment",
        "comparator_alignment",
        "intervention_or_exposure_alignment",
        "direction_alignment",
    ],
)
def test_v3_alignment_mismatch_prevents_support(field):
    result = judgment()
    result.question_set_version = "claim-source-en-v3"
    for name in (
        "outcome_alignment",
        "timeframe_alignment",
        "comparator_alignment",
        "intervention_or_exposure_alignment",
        "direction_alignment",
    ):
        setattr(result, name, choice("aligned"))
    setattr(result, field, choice("mismatch"))
    policy = propose_disposition(
        run_deterministic_checks("Claim.", "Claim."),
        result,
        source_access_tier=SourceAccessTier.FULL_TEXT,
    )
    assert policy.status == ProposalStatus.PARTIALLY_SUPPORTED
    assert not policy.auto_accepted
    assert any(field.split("_")[0] in r for r in policy.reasons)
    setattr(result, field, choice("unclear"))
    assert (
        propose_disposition(
            run_deterministic_checks("Claim.", "Claim."),
            result,
            source_access_tier=SourceAccessTier.FULL_TEXT,
        ).status
        == ProposalStatus.REVIEW_REQUIRED
    )


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "mutation", ["missing", "unknown-option", "nan", "bad-distribution", "root-list"]
)
async def test_malformed_v3_response_fails_closed(mutation):
    body = response_body()
    if mutation == "missing":
        del body["answers"]["outcome_alignment"]
    elif mutation == "unknown-option":
        body["answers"]["outcome_alignment"]["choice"] = "invented"
    elif mutation == "nan":
        body["answers"]["outcome_alignment"]["probabilities"]["aligned"] = float("nan")
    elif mutation == "bad-distribution":
        body["answers"]["outcome_alignment"]["probabilities"]["aligned"] = 0.1
    elif mutation == "root-list":
        body = []

    async def handler(request):
        if mutation == "nan":
            return httpx.Response(
                200, content=json.dumps(body), headers={"content-type": "application/json"}
            )
        return httpx.Response(200, json=body)

    provider = TypeSafeJevProvider(
        api_key="test-only",
        endpoint="https://example.test",
        model="jev-1.13.0",
        transport=httpx.MockTransport(handler),
    )
    with pytest.raises(ProviderError):
        await provider.evaluate("Claim.", "Source evidence.")


@pytest.mark.parametrize("failure", [False, True])
def test_evidence_correction_is_atomic_versioned_and_idempotent(tmp_path, failure):
    provider = RevisionProvider(error=ProviderError("test failure") if failure else None)
    app = create_app(Settings(data_dir=tmp_path, jev_api_key=None))
    app.state.judgment_provider = provider
    with TestClient(app) as client:
        audit = client.post(
            "/api/v1/audits",
            json={
                "claim": "The intervention decreased the endpoint.",
                "source_text": source_text(),
            },
        ).json()
        base = f"/api/v1/audits/{audit['audit_id']}"
        original = client.get(base + "/proposals").json()[0]
        selected = [audit["candidates"][-1]["passage"]["passage_id"]]
        request = {
            **reference(audit),
            "reviewer": "human",
            "feedback": "Inspect the other cohort.",
            "idempotency_key": "evidence-test-key",
            "selected_passage_ids": selected,
        }
        invalid = {**request, "selected_passage_ids": ["foreign-source-id"]}
        assert client.post(base + "/revisions", json=invalid).status_code == 409
        assert len(provider.calls) == 1
        response = client.post(base + "/revisions", json=request)
        assert response.status_code == 200
        assert response.json()["status"] == ("revision_failed" if failure else "revision_completed")
        current = client.get(base).json()
        assert client.get(base + "/proposals").json()[0]["evidence_set"] == original["evidence_set"]
        assert client.post(base + "/revisions", json=request).json() == response.json()
        assert len(provider.calls) == 2
        if failure:
            assert current["evidence_set"] == audit["evidence_set"]
            assert current["current_proposal_version"] == 1
        else:
            assert current["current_proposal_version"] == 2
            assert [p["passage_id"] for p in current["evidence_set"]["passages"]] == selected
            events = client.get(base + "/events").json()
            change = next(e for e in events if e["event_type"] == "evidence.selection.changed")
            assert change["payload"]["reviewer"] == "human"
            assert change["payload"]["new_hash"] == current["provenance"]["evidence_set_hash"]
            assert change["proposal_id"] == current["current_proposal_id"]
            assert change["proposal_version"] == 2
            assert provider.calls[-1][1] == evidence_text(
                make_evidence_set([app.state.store.get(audit["audit_id"]).selected_passage])
            )
            assert (
                client.post(
                    base + "/reviews",
                    json={
                        **reference(audit),
                        "decision": "accept",
                        "reviewer": "human",
                        "notes": "stale",
                    },
                ).status_code
                == 409
            )


def test_local_evidence_correction_without_provider_remains_review_required(tmp_path):
    app = create_app(Settings(data_dir=tmp_path, jev_api_key=None))
    with TestClient(app) as client:
        audit = client.post(
            "/api/v1/audits",
            json={
                "claim": "The intervention decreased the endpoint.",
                "source_text": source_text(),
                "use_judgment_provider": False,
            },
        ).json()
        base = f"/api/v1/audits/{audit['audit_id']}"
        result = client.post(
            base + "/revisions",
            json={
                **reference(audit),
                "reviewer": "human",
                "feedback": "Select the correct paragraph.",
                "idempotency_key": "local-evidence-key",
                "selected_passage_ids": [audit["candidates"][-1]["passage"]["passage_id"]],
            },
        )
        assert result.json()["status"] == "revision_completed"
        current = client.get(base).json()
        assert current["proposal"]["status"] == "review_required"
        assert current["judgment_result"] is None
        assert not current["proposal"]["auto_accepted"]
