from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from claim_trellis.citations import extract_citation_sentences
from claim_trellis.document_blocks import PARSER_VERSION, make_block, text_blocks, validate_blocks
from claim_trellis.models import DocumentBlock, ParsedDocument


class IngestionError(ValueError):
    pass


SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _read_pdf(data: bytes) -> tuple[str, list[str], list[tuple[str, dict[str, str | int]]]]:
    if not data.startswith(b"%PDF"):
        raise IngestionError("The file extension is PDF but the signature is not a PDF.")
    try:
        reader = PdfReader(io.BytesIO(data), strict=True)
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # pypdf raises several parser-specific exception classes
        raise IngestionError(f"PDF parsing failed: {exc}") from exc
    text = "\n\n".join(pages)
    warnings: list[str] = []
    if not text.strip():
        warnings.append("No selectable text was found; scanned PDFs require OCR before upload.")
    return text, warnings, [(page, {"page": index}) for index, page in enumerate(pages, 1)]


def _read_docx(data: bytes) -> tuple[str, list[str], list[tuple[str, dict[str, str | int]]]]:
    if not data.startswith(b"PK"):
        raise IngestionError(
            "The file extension is DOCX but the signature is not an Office archive."
        )
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestionError(f"DOCX parsing failed: {exc}") from exc
    blocks = [paragraph.text for paragraph in document.paragraphs]
    descriptors: list[tuple[str, dict[str, str | int]]] = []
    section = ""
    for index, paragraph in enumerate(document.paragraphs, 1):
        style = paragraph.style.name if paragraph.style else ""
        is_heading = style.startswith("Heading")
        if is_heading:
            section = paragraph.text
        descriptors.append(
            (
                paragraph.text,
                {
                    "section": section,
                    "paragraph": index,
                    "block_type": "heading" if is_heading else "paragraph",
                },
            )
        )
    for table in document.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text for cell in row.cells)
            blocks.append(row_text)
            descriptors.append(
                (row_text, {"section": "", "paragraph": len(blocks), "block_type": "table_row"})
            )
    return "\n\n".join(blocks), [], descriptors


def parse_document_bytes(
    filename: str,
    data: bytes,
    media_type: str = "application/octet-stream",
    *,
    max_bytes: int = 25 * 1024 * 1024,
    max_chars: int = 2_000_000,
) -> ParsedDocument:
    if len(data) > max_bytes:
        raise IngestionError(f"Upload exceeds the {max_bytes:,}-byte limit.")
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise IngestionError(
            f"Unsupported file type {suffix or '(none)'}. Use TXT, MD, PDF, or DOCX."
        )

    warnings: list[str] = []
    descriptors: list[tuple[str, dict[str, str | int]]] = []
    if suffix == ".pdf":
        raw_text, warnings, descriptors = _read_pdf(data)
    elif suffix == ".docx":
        raw_text, warnings, descriptors = _read_docx(data)
    else:
        try:
            raw_text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise IngestionError("Text files must use UTF-8 encoding.") from exc

    text = _clean_text(raw_text)
    if len(text) > max_chars:
        raise IngestionError(f"Parsed text exceeds the {max_chars:,}-character limit.")
    if not text:
        raise IngestionError("The document contains no readable text.")

    structured: list[DocumentBlock] = []
    cursor = 0
    for content, descriptor in descriptors:
        normalized = _clean_text(content)
        if not normalized:
            continue
        start = text.find(normalized, cursor)
        if start < 0:
            warnings.append("A format locator could not be grounded; normalized offsets retained.")
            continue
        end = start + len(normalized)
        page = int(descriptor["page"]) if "page" in descriptor else None
        paragraph = int(descriptor["paragraph"]) if "paragraph" in descriptor else None
        section = str(descriptor.get("section", "")) or None
        structured.append(
            make_block(
                text,
                start,
                end,
                len(structured) + 1,
                page=page,
                section=section,
                paragraph=paragraph,
                block_type=str(descriptor.get("block_type", "paragraph")),
                locator=f"page:{page} / chars:{start}-{end}"
                if page
                else f"section:{section or '(none)'} / paragraph:{paragraph}",
            )
        )
        cursor = end
    if not descriptors:
        structured = text_blocks(text, markdown=suffix == ".md")
    else:
        try:
            validate_blocks(text, structured)
        except ValueError:
            # Never omit source text or invent a page/section when a format locator fails.
            structured = text_blocks(text)
            warnings.append(
                "Format locators could not cover the normalized text; using exact text offsets."
            )

    return ParsedDocument(
        filename=Path(filename).name,
        media_type=media_type,
        text=text,
        content_sha256=hashlib.sha256(text.encode()).hexdigest(),
        character_count=len(text),
        citation_sentences=extract_citation_sentences(text),
        warnings=warnings,
        blocks=structured,
        parser_version=PARSER_VERSION,
    )
